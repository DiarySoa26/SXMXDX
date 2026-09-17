import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer


# ============================================================
# CONFIGURATION
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

DEFAULT_EXCEL_FILE = (
    PROJECT_ROOT
    / "data"
    / "reference"
    / "BP_SIMULATION_SOMIDA_2024.xlsm"
)

OUTPUT_DIR = (
    PROJECT_ROOT
    / "reverse_engineering"
    / "output"
)


def clean_value(value):
    """
    Transforme une valeur Excel en valeur sérialisable.
    """
    if value is None:
        return ""

    return str(value).replace("\n", " ").strip()


def classify_cell(cell):
    """
    Classification technique initiale d'une cellule.
    """

    if cell.value is None:
        return "EMPTY"

    if cell.data_type == "f":
        return "FORMULA"

    if isinstance(cell.value, str):
        return "TEXT"

    if isinstance(cell.value, bool):
        return "BOOLEAN"

    if isinstance(cell.value, (int, float)):
        return "NUMBER"

    return "OTHER"


def normalize_sheet_name(sheet_name):
    if sheet_name is None:
        return ""

    return sheet_name.replace("''", "'").strip("'")



def extract_formula_references(formula, current_sheet):
    """
    Analyse une formule Excel avec le tokenizer openpyxl.
    """

    references = []

    if not formula or not isinstance(formula, str):
        return references

    try:
        tokenizer = Tokenizer(formula)
    except Exception:
        return references

    for token in tokenizer.items:

        if token.type != "OPERAND":
            continue

        if token.subtype != "RANGE":
            continue

        value = token.value

        if not value:
            continue

        # Référence externe : [Workbook.xlsx]Sheet!A1
        if "[" in value and "]" in value:
            references.append({
                "target_sheet": "",
                "target_reference": value,
                "reference_type": "EXTERNAL"
            })
            continue

        # Référence inter-feuille
        if "!" in value:

            sheet_part, reference_part = value.rsplit("!", 1)

            target_sheet = normalize_sheet_name(sheet_part)

            references.append({
                "target_sheet": target_sheet,
                "target_reference": reference_part,
                "reference_type": (
                    "RANGE"
                    if ":" in reference_part
                    else "CELL"
                )
            })

        else:

            # A1 / A1:B12
            cell_reference_pattern = (
                r"^\$?[A-Z]{1,3}\$?\d+"
                r"(?::\$?[A-Z]{1,3}\$?\d+)?$"
            )

            if re.match(cell_reference_pattern, value):

                references.append({
                    "target_sheet": current_sheet,
                    "target_reference": value,
                    "reference_type": (
                        "RANGE"
                        if ":" in value
                        else "CELL"
                    )
                })

            else:

                # Peut être un Named Range Excel
                references.append({
                    "target_sheet": "",
                    "target_reference": value,
                    "reference_type": "NAME"
                })

    return references


def write_csv(path, fieldnames, rows):

    with path.open(
        "w",
        newline="",
        encoding="utf-8-sig"
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames
        )

        writer.writeheader()
        writer.writerows(rows)


def inspect_workbook(excel_file):

    print("=" * 70)
    print("SXMXDX - REVERSE ENGINEERING EXCEL")
    print("=" * 70)

    print(f"\nFichier : {excel_file}")

    if not excel_file.exists():
        raise FileNotFoundError(
            f"Fichier introuvable : {excel_file}"
        )

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    print("\nChargement du classeur...")

    workbook = load_workbook(
        excel_file,
        data_only=False,
        read_only=False,
        keep_vba=True,
        keep_links=True
    )

    print(
        f"Classeur chargé : "
        f"{len(workbook.sheetnames)} feuilles."
    )


    sheets_rows = []
    cells_rows = []
    formulas_rows = []
    dependencies_rows = []
    external_rows = []

    sheet_dependency_counter = Counter()

    formulas_by_sheet = Counter()
    constants_by_sheet = Counter()

    global_stats = {
        "workbook": excel_file.name,
        "sheet_count": len(workbook.sheetnames),
        "non_empty_cells": 0,
        "formula_cells": 0,
        "constant_cells": 0,
        "external_references": 0
    }


    # Parcours des feuilles

    for index, worksheet in enumerate(
        workbook.worksheets,
        start=1
    ):

        print(
            f"[{index:02d}/{len(workbook.worksheets)}] "
            f"Analyse : {worksheet.title}"
        )

        sheet_non_empty = 0
        sheet_formulas = 0
        sheet_constants = 0

        min_row_used = None
        max_row_used = None
        min_col_used = None
        max_col_used = None


        # Parcours des cellules
  

        for row in worksheet.iter_rows():

            for cell in row:

                if cell.value is None:
                    continue

                sheet_non_empty += 1
                global_stats["non_empty_cells"] += 1

                # Dimensions réellement utilisées
                if min_row_used is None:
                    min_row_used = cell.row
                    max_row_used = cell.row
                    min_col_used = cell.column
                    max_col_used = cell.column
                else:
                    min_row_used = min(
                        min_row_used,
                        cell.row
                    )

                    max_row_used = max(
                        max_row_used,
                        cell.row
                    )

                    min_col_used = min(
                        min_col_used,
                        cell.column
                    )

                    max_col_used = max(
                        max_col_used,
                        cell.column
                    )

                cell_type = classify_cell(cell)


                cells_rows.append({
                    "sheet": worksheet.title,
                    "cell": cell.coordinate,
                    "row": cell.row,
                    "column": cell.column,
                    "type": cell_type,
                    "value": clean_value(cell.value),
                    "number_format": cell.number_format,
                    "style_id": cell.style_id
                })


                if cell_type == "FORMULA":

                    sheet_formulas += 1

                    global_stats["formula_cells"] += 1

                    formulas_by_sheet[
                        worksheet.title
                    ] += 1

                    formula = str(cell.value)

                    references = (
                        extract_formula_references(
                            formula,
                            worksheet.title
                        )
                    )

                    formulas_rows.append({
                        "sheet": worksheet.title,
                        "cell": cell.coordinate,
                        "formula": formula,
                        "reference_count": len(references)
                    })

                    # Dépendances

                    for reference in references:

                        target_sheet = (
                            reference["target_sheet"]
                        )

                        target_reference = (
                            reference["target_reference"]
                        )

                        reference_type = (
                            reference["reference_type"]
                        )

                        dependencies_rows.append({
                            "source_sheet":
                                worksheet.title,

                            "source_cell":
                                cell.coordinate,

                            "target_sheet":
                                target_sheet,

                            "target_reference":
                                target_reference,

                            "reference_type":
                                reference_type
                        })

                        if (
                            reference_type
                            == "EXTERNAL"
                        ):

                            external_rows.append({
                                "source_sheet":
                                    worksheet.title,

                                "source_cell":
                                    cell.coordinate,

                                "formula":
                                    formula,

                                "external_reference":
                                    target_reference
                            })

                            global_stats[
                                "external_references"
                            ] += 1

                        elif target_sheet:

                            sheet_dependency_counter[
                                (
                                    worksheet.title,
                                    target_sheet
                                )
                            ] += 1

                else:

                    sheet_constants += 1

                    global_stats[
                        "constant_cells"
                    ] += 1

                    constants_by_sheet[
                        worksheet.title
                    ] += 1


        # Informations feuille

        if min_row_used is None:
            used_range = ""
        else:
            used_range = (
                f"{worksheet.cell(min_row_used, min_col_used).coordinate}"
                f":"
                f"{worksheet.cell(max_row_used, max_col_used).coordinate}"
            )

        formula_ratio = 0

        if sheet_non_empty > 0:
            formula_ratio = (
                sheet_formulas
                / sheet_non_empty
            )

        sheets_rows.append({
            "sheet_index": index,
            "sheet": worksheet.title,
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "used_range": used_range,
            "non_empty_cells": sheet_non_empty,
            "constant_cells": sheet_constants,
            "formula_cells": sheet_formulas,
            "formula_ratio": round(
                formula_ratio,
                4
            )
        })

    # DEPENDANCES ENTRE FEUILLES

    sheet_dependencies_rows = []

    for (
        source_sheet,
        target_sheet
    ), count in sorted(
        sheet_dependency_counter.items()
    ):

        sheet_dependencies_rows.append({
            "source_sheet": source_sheet,
            "target_sheet": target_sheet,
            "reference_count": count,
            "self_reference":
                source_sheet == target_sheet
        })

    print("\nÉcriture des résultats...")

    write_csv(
        OUTPUT_DIR / "sheets.csv",
        [
            "sheet_index",
            "sheet",
            "max_row",
            "max_column",
            "used_range",
            "non_empty_cells",
            "constant_cells",
            "formula_cells",
            "formula_ratio"
        ],
        sheets_rows
    )

    write_csv(
        OUTPUT_DIR / "cells.csv",
        [
            "sheet",
            "cell",
            "row",
            "column",
            "type",
            "value",
            "number_format",
            "style_id"
        ],
        cells_rows
    )

    write_csv(
        OUTPUT_DIR / "formulas.csv",
        [
            "sheet",
            "cell",
            "formula",
            "reference_count"
        ],
        formulas_rows
    )

    write_csv(
        OUTPUT_DIR / "dependencies.csv",
        [
            "source_sheet",
            "source_cell",
            "target_sheet",
            "target_reference",
            "reference_type"
        ],
        dependencies_rows
    )

    write_csv(
        OUTPUT_DIR / "sheet_dependencies.csv",
        [
            "source_sheet",
            "target_sheet",
            "reference_count",
            "self_reference"
        ],
        sheet_dependencies_rows
    )

    write_csv(
        OUTPUT_DIR / "external_references.csv",
        [
            "source_sheet",
            "source_cell",
            "formula",
            "external_reference"
        ],
        external_rows
    )

    # JSON SUMMARY
    summary = {
        "global": global_stats,
        "sheets": sheets_rows
    }

    with (
        OUTPUT_DIR / "summary.json"
    ).open(
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            summary,
            file,
            indent=4,
            ensure_ascii=False
        )

    print("\n" + "=" * 70)
    print("ANALYSE TERMINÉE")
    print("=" * 70)

    print(
        f"Feuilles             : "
        f"{global_stats['sheet_count']}"
    )

    print(
        f"Cellules non vides   : "
        f"{global_stats['non_empty_cells']}"
    )

    print(
        f"Formules             : "
        f"{global_stats['formula_cells']}"
    )

    print(
        f"Constantes           : "
        f"{global_stats['constant_cells']}"
    )

    print(
        f"Références externes  : "
        f"{global_stats['external_references']}"
    )

    print(
        f"\nRésultats disponibles dans :\n"
        f"{OUTPUT_DIR}"
    )

    print("\nFichiers générés :")

    for filename in [
        "sheets.csv",
        "cells.csv",
        "formulas.csv",
        "dependencies.csv",
        "sheet_dependencies.csv",
        "external_references.csv",
        "summary.json"
    ]:
        print(f" - {filename}")



if __name__ == "__main__":

    if len(sys.argv) > 1:
        excel_path = Path(sys.argv[1]).resolve()
    else:
        excel_path = DEFAULT_EXCEL_FILE

    inspect_workbook(excel_path)