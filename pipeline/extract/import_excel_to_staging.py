import hashlib
import os
import sys
from datetime import datetime
from pathlib import Path

import psycopg
from dotenv import load_dotenv
from openpyxl import load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

load_dotenv(PROJECT_ROOT / ".env")


def get_connection():

    return psycopg.connect(
        host=os.getenv("POSTGRES_HOST"),
        port=os.getenv("POSTGRES_PORT"),
        dbname=os.getenv("POSTGRES_DB"),
        user=os.getenv("POSTGRES_USER"),
        password=os.getenv("POSTGRES_PASSWORD"),
    )


# ============================================================
# CHECKSUM
# ============================================================

def calculate_checksum(file_path):

    sha256 = hashlib.sha256()

    with file_path.open("rb") as file:

        for block in iter(
            lambda: file.read(1024 * 1024),
            b""
        ):
            sha256.update(block)

    return sha256.hexdigest()


# ============================================================
# EXERCICE
# ============================================================

def get_or_create_exercice(
    connection,
    year,
    filename
):

    with connection.cursor() as cursor:

        cursor.execute(
            """
            SELECT exercice_id
            FROM sxmxdx.exercice
            WHERE annee = %s
            """,
            (year,)
        )

        result = cursor.fetchone()

        if result:
            return result[0]

        cursor.execute(
            """
            INSERT INTO sxmxdx.exercice
            (
                annee,
                libelle,
                statut,
                fichier_source
            )
            VALUES
            (
                %s,
                %s,
                'REFERENCE',
                %s
            )
            RETURNING exercice_id
            """,
            (
                year,
                f"Budget SOMIDA {year}",
                filename
            )
        )

        exercice_id = cursor.fetchone()[0]

        connection.commit()

        return exercice_id


# ============================================================
# IMPORT BATCH
# ============================================================

def create_import_batch(
    connection,
    exercice_id,
    filename,
    checksum
):

    with connection.cursor() as cursor:

        cursor.execute(
            """
            INSERT INTO sxmxdx.import_batch
            (
                exercice_id,
                fichier,
                checksum,
                statut,
                date_debut
            )
            VALUES
            (
                %s,
                %s,
                %s,
                'RUNNING',
                CURRENT_TIMESTAMP
            )
            RETURNING import_batch_id
            """,
            (
                exercice_id,
                filename,
                checksum
            )
        )

        import_batch_id = (
            cursor.fetchone()[0]
        )

        connection.commit()

        return import_batch_id


# ============================================================
# CLASSIFICATION TECHNIQUE
# ============================================================

def get_technical_type(cell):

    if cell.data_type == "f":
        return "FORMULA"

    value = cell.value

    if isinstance(value, bool):
        return "BOOLEAN"

    if isinstance(value, (int, float)):
        return "NUMBER"

    if isinstance(value, datetime):
        return "DATE"

    if isinstance(value, str):
        return "TEXT"

    return "OTHER"


# ============================================================
# INSERTION STAGING
# ============================================================

def insert_cells(
    connection,
    workbook,
    import_batch_id
):

    sql = """
        INSERT INTO staging.import_cell
        (
            import_batch_id,
            sheet_name,
            cell_reference,
            row_number,
            column_number,
            technical_type,
            raw_value,
            number_format
        )
        VALUES
        (
            %s, %s, %s, %s,
            %s, %s, %s, %s
        )
    """

    batch = []

    total = 0

    BATCH_SIZE = 2000

    with connection.cursor() as cursor:

        for worksheet in workbook.worksheets:

            print(
                f"Import : {worksheet.title}"
            )

            sheet_count = 0

            for row in worksheet.iter_rows():

                for cell in row:

                    if cell.value is None:
                        continue

                    technical_type = (
                        get_technical_type(cell)
                    )

                    raw_value = str(cell.value)

                    batch.append(
                        (
                            import_batch_id,
                            worksheet.title,
                            cell.coordinate,
                            cell.row,
                            cell.column,
                            technical_type,
                            raw_value,
                            cell.number_format
                        )
                    )

                    total += 1
                    sheet_count += 1

                    if len(batch) >= BATCH_SIZE:

                        cursor.executemany(
                            sql,
                            batch
                        )

                        connection.commit()

                        batch.clear()

            print(
                f"  -> {sheet_count} cellules"
            )

        if batch:

            cursor.executemany(
                sql,
                batch
            )

            connection.commit()

    return total


# ============================================================
# FINALISATION
# ============================================================

def finish_import(
    connection,
    import_batch_id,
    total,
    status="IMPORTED"
):

    with connection.cursor() as cursor:

        cursor.execute(
            """
            UPDATE sxmxdx.import_batch
            SET
                statut = %s,
                nombre_lignes = %s,
                date_fin = CURRENT_TIMESTAMP
            WHERE import_batch_id = %s
            """,
            (
                status,
                total,
                import_batch_id
            )
        )

        connection.commit()


def fail_import(
    connection,
    import_batch_id
):

    with connection.cursor() as cursor:

        cursor.execute(
            """
            UPDATE sxmxdx.import_batch
            SET
                statut = 'FAILED',
                date_fin = CURRENT_TIMESTAMP
            WHERE import_batch_id = %s
            """,
            (import_batch_id,)
        )

        connection.commit()


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 72)
    print("SXMXDX - IMPORT EXCEL -> POSTGRESQL STAGING")
    print("=" * 72)

    # --------------------------------------------------------
    # Fichier
    # --------------------------------------------------------

    if len(sys.argv) > 1:

        excel_file = Path(
            sys.argv[1]
        ).resolve()

    else:

        configured_file = os.getenv(
            "REFERENCE_EXCEL_FILE"
        )

        excel_file = (
            PROJECT_ROOT
            / configured_file
        ).resolve()

    if not excel_file.exists():

        raise FileNotFoundError(
            f"Fichier absent : {excel_file}"
        )

    # --------------------------------------------------------
    # Pour le modèle canonique actuel
    # --------------------------------------------------------

    year = 2024

    print(f"\nFichier  : {excel_file.name}")
    print(f"Exercice : {year}")

    checksum = calculate_checksum(
        excel_file
    )

    print(
        f"SHA256   : {checksum[:16]}..."
    )

    connection = get_connection()

    import_batch_id = None

    try:

        exercice_id = (
            get_or_create_exercice(
                connection,
                year,
                excel_file.name
            )
        )

        print(
            f"exercice_id     : "
            f"{exercice_id}"
        )

        import_batch_id = (
            create_import_batch(
                connection,
                exercice_id,
                excel_file.name,
                checksum
            )
        )

        print(
            f"import_batch_id : "
            f"{import_batch_id}"
        )

        print(
            "\nChargement Excel..."
        )

        workbook = load_workbook(
            excel_file,
            data_only=False,
            read_only=False,
            keep_vba=True,
            keep_links=True
        )

        print(
            f"{len(workbook.sheetnames)} "
            f"feuilles détectées.\n"
        )

        total = insert_cells(
            connection,
            workbook,
            import_batch_id
        )

        finish_import(
            connection,
            import_batch_id,
            total
        )

        print("\n" + "=" * 72)
        print("IMPORT TERMINÉ")
        print("=" * 72)

        print(
            f"\nCellules importées : "
            f"{total}"
        )

        print(
            f"Batch : "
            f"{import_batch_id}"
        )

    except Exception:

        if import_batch_id is not None:

            fail_import(
                connection,
                import_batch_id
            )

        raise

    finally:

        connection.close()


if __name__ == "__main__":
    main()