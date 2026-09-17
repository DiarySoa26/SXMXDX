import openpyxl
import re

EXCEL_FILE = "BP_SIMULATION_SOMIDA_2024.xlsm"
TARGET_SHEET = "OKB TRA ABR"

wb = openpyxl.load_workbook(
    EXCEL_FILE,
    data_only=False,
    keep_vba=True
)

# ws = wb[TARGET_SHEET]

print("\n=== NOMS EXACTS DES FEUILLES ===")

for i, sheet_name in enumerate(wb.sheetnames, 1):
    print(f"{i:02d}. {repr(sheet_name)}")

exit()

print("=" * 80)
print(f"FEUILLE : {TARGET_SHEET}")
print("=" * 80)

print(f"Dimensions : {ws.max_row} lignes x {ws.max_column} colonnes")

# --------------------------------------------------
# 1. Afficher les cellules non vides
# --------------------------------------------------

print("\n--- CONTENU ---")

count = 0

for row in ws.iter_rows():
    values = []

    for cell in row:
        if cell.value is not None:
            values.append(
                f"{cell.coordinate}={str(cell.value)[:100]}"
            )

    if values:
        print(" | ".join(values))
        count += 1

    # Eviter une sortie gigantesque
    if count >= 40:
        print("...")
        break


# --------------------------------------------------
# 2. Compter les formules
# --------------------------------------------------

formulas = []

for row in ws.iter_rows():
    for cell in row:

        if (
            isinstance(cell.value, str)
            and cell.value.startswith("=")
        ):
            formulas.append(
                (cell.coordinate, cell.value)
            )

print(f"\n--- FORMULES : {len(formulas)} ---")

for coordinate, formula in formulas[:30]:
    print(f"{coordinate} -> {formula}")

if len(formulas) > 30:
    print("...")


# --------------------------------------------------
# 3. Chercher les feuilles référencées
#    PAR OKB TRA ABR
# --------------------------------------------------

referenced_sheets = set()

for _, formula in formulas:

    # références de type 'Nom feuille'!A1
    matches = re.findall(
        r"'([^']+)'!",
        formula
    )

    referenced_sheets.update(matches)

print("\n--- FEUILLES UTILISÉES PAR OKB TRA ABR ---")

if referenced_sheets:
    for sheet in sorted(referenced_sheets):
        print(" ->", sheet)
else:
    print("Aucune référence explicite trouvée.")


# --------------------------------------------------
# 4. Chercher les feuilles qui utilisent
#    OKB TRA ABR
# --------------------------------------------------

used_by = []

for sheet_name in wb.sheetnames:

    if sheet_name == TARGET_SHEET:
        continue

    other_ws = wb[sheet_name]

    for row in other_ws.iter_rows():
        for cell in row:

            value = cell.value

            if (
                isinstance(value, str)
                and value.startswith("=")
                and TARGET_SHEET.lower() in value.lower()
            ):
                used_by.append(
                    (
                        sheet_name,
                        cell.coordinate,
                        value
                    )
                )

print("\n--- OKB TRA ABR EST UTILISÉE PAR ---")

if used_by:

    for sheet, coordinate, formula in used_by[:50]:

        print(
            f"{sheet}!{coordinate}"
            f" -> {formula}"
        )

else:
    print("Aucune autre feuille ne semble la référencer.")


print("\nAnalyse terminée.")